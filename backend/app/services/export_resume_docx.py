"""ATS-friendly DOCX export of a tailored resume (python-docx).

Same rules as the PDF: one column, no tables / text boxes / images / headers / footers, standard
section headings, real bullet paragraphs (Word's "List Bullet" style, not typed characters), and
common fonts. It uses the same template keys as the PDF so the two files look alike.
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.services.export_resume_pdf import DEFAULT_TEMPLATE, TEMPLATES, _BULLET_RE, Template

_PAGES_MM = {"a4": (210, 297), "letter": (215.9, 279.4)}
_WORD_FONT = {"Helvetica": "Arial", "Times-Roman": "Times New Roman"}

# Characters Word/XML cannot hold (control chars); everything else, including Unicode, is fine in DOCX.
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _t(s: Any) -> str:
    return _BAD_XML.sub("", str(s or ""))


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _shade(paragraph, hex_color: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    ppr.append(shd)


def _style_font(style, name: str, size: float, bold: bool = False, color: str = "#000000") -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = _rgb(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _bottom_border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "000000")
    bdr.append(b)
    ppr.append(bdr)


def _spacing(p, before: float = 0, after: float = 0, line: float | None = None) -> None:
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(before), Pt(after)
    if line:
        f.line_spacing = Pt(line)


def render_resume_docx(resume: dict[str, Any], *, template: str = DEFAULT_TEMPLATE, page: str = "a4") -> bytes:
    t: Template = TEMPLATES.get(template) or TEMPLATES[DEFAULT_TEMPLATE]
    font = _WORD_FONT.get(t.font, "Arial")
    doc = Document()

    w, h = _PAGES_MM.get(page, _PAGES_MM["a4"])
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(w), Mm(h)
    sec.left_margin = sec.right_margin = Mm(t.margin_mm)
    sec.top_margin = sec.bottom_margin = Mm(t.margin_mm * 0.9)

    _style_font(doc.styles["Normal"], font, t.body)
    _style_font(doc.styles["List Bullet"], font, t.body)
    _style_font(doc.styles["Heading 1"], font, t.head_size, bold=True, color=t.accent)   # real heading style: parsers use it
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True

    p = resume.get("profile") or {}
    name = _t(p.get("name")).strip() or "Candidate"

    # name
    para = doc.add_paragraph()
    r = para.add_run(name)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(t.name_size), _rgb(t.accent)
    if t.align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(para, 0, 2)

    contact = "  |  ".join(_t(c) for c in [p.get("email"), p.get("phone"), p.get("city"), p.get("github")] if c and str(c).strip())
    if contact:
        para = doc.add_paragraph(contact)
        if t.align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(para, 0, 2)

    def heading(text: str) -> None:
        hp = doc.add_paragraph(text.upper(), style="Heading 1")
        _spacing(hp, t.section_gap + (3 if t.head_bar else 0), 3)
        if t.head_bar:
            _shade(hp, t.bar_color)
        elif t.rule:
            _bottom_border(hp)

    def body_lines(content: Any) -> None:
        for line in str(content or "").splitlines():
            if not line.strip():
                continue
            if _BULLET_RE.match(line):
                bp = doc.add_paragraph(_t(_BULLET_RE.sub("", line)), style="List Bullet")
                _spacing(bp, 0, 1)
            else:
                bp = doc.add_paragraph(_t(line.strip()))
                _spacing(bp, 0, 1)

    def title_line(text: str) -> None:
        tp = doc.add_paragraph()
        tp.add_run(_t(text)).bold = True
        _spacing(tp, 4, 0)
        tp.paragraph_format.keep_with_next = True

    if str(resume.get("summary") or "").strip():
        heading("Summary")
        doc.add_paragraph(_t(resume["summary"]))

    skills = [str(s).strip() for s in resume.get("skills") or [] if str(s).strip()]
    if skills:
        heading("Skills")
        doc.add_paragraph(_t(", ".join(skills)))

    def entries(key: str, label: str) -> None:
        items = [it for it in resume.get(key) or [] if it]
        rendered = []
        for it in items:
            if isinstance(it, dict):
                rendered.append((it.get("title") or "", it.get("content") or ""))
            else:
                rendered.append(("", str(it)))
        rendered = [(a, b) for a, b in rendered if str(a).strip() or str(b).strip()]
        if not rendered:
            return
        heading(label)
        for title, content in rendered:
            if str(title).strip():
                title_line(title)
            body_lines(content)

    entries("experience", "Experience")
    entries("projects", "Projects")

    edu = [e for e in resume.get("education") or [] if isinstance(e, dict) and (e.get("school") or e.get("degree") or e.get("major"))]
    if edu:
        heading("Education")
        for e in edu:
            if e.get("school"):
                title_line(e["school"])
            degree = " ".join(x for x in [e.get("degree"), e.get("major")] if x)
            dates = " - ".join(x for x in [e.get("start_date"), e.get("end_date")] if x)
            line = ", ".join(x for x in [degree, dates] if x)
            if line:
                doc.add_paragraph(_t(line))

    entries("awards", "Awards")

    cp = doc.core_properties
    cp.title, cp.author, cp.subject = f"{name} - Resume", name, "Resume"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
