"""Edit → approve → generate & validate: DOCX export, file validation, re-scoring, readiness verdict."""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.services.ats_score import compare, rescore, unverified_skills
from app.services.export_resume_docx import render_resume_docx
from app.services.export_resume_pdf import TEMPLATES, render_resume_pdf
from app.services.readiness import verdict
from app.services.resume_validation import expected_sections, validate_docx, validate_pdf
from test_resume_export_and_match import JD, ORIGINAL, TAILORED


# ------------------------------------------------------------------ DOCX export
@pytest.mark.parametrize("template", list(TEMPLATES))
def test_docx_is_ats_safe(template):
    data = render_resume_docx(TAILORED, template=template)
    doc = Document(io.BytesIO(data))
    assert not doc.tables and not doc.inline_shapes
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 1"] == [
        "SUMMARY", "SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION", "AWARDS"]
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert "Designed REST APIs serving 2M req/day & cut latency by 40%" in bullets   # real bullets, not "-" text
    assert doc.core_properties.title == "Rahul Sharma - Resume"


def test_docx_keeps_unicode_that_the_pdf_fonts_cannot():
    r = {"profile": {"name": "Zoë Ñandú"}, "summary": "Earned ₹12L — “top” performer"}
    assert "₹12L" in "\n".join(p.text for p in Document(io.BytesIO(render_resume_docx(r))).paragraphs)


# ------------------------------------------------------------------ validation: good files pass
@pytest.mark.parametrize("template", list(TEMPLATES))
def test_generated_files_validate(template):
    pdf = validate_pdf(render_resume_pdf(TAILORED, template=template), TAILORED)
    docx = validate_docx(render_resume_docx(TAILORED, template=template), TAILORED)
    assert pdf["ok"], pdf
    assert docx["ok"], docx
    assert pdf["pages"] == 1
    assert {c["id"] for c in pdf["checks"]} >= {"opens", "text", "sections", "contact", "content", "length"}


def test_validation_handles_unicode_resume():
    r = {"profile": {"name": "Zoë Ñandú", "email": "z@x.com"}, "summary": "Earned ₹12L — “top” performer",
         "skills": ["Python"], "experience": [{"title": "Dev", "content": "• shipped café ordering"}]}
    assert validate_pdf(render_resume_pdf(r), r)["ok"]
    assert validate_docx(render_resume_docx(r), r)["ok"]


def test_expected_sections_skips_empty():
    assert expected_sections({"summary": "x", "skills": [], "experience": [], "education": [{"school": "MIT"}]}) == ["SUMMARY", "EDUCATION"]


# ------------------------------------------------------------------ validation: bad files FAIL
def test_corrupt_pdf_fails():
    v = validate_pdf(b"%PDF-1.4 not really a pdf", TAILORED)
    assert not v["ok"] and v["checks"][0]["status"] == "fail"


def test_pdf_missing_content_fails():
    other = {"profile": {"name": "Someone Else", "email": "else@example.org"}, "summary": "Totally different words here",
             "experience": [{"title": "Other", "content": "- unrelated things happened"}]}
    v = validate_pdf(render_resume_pdf(other), TAILORED)      # file does not match the resume it should contain
    assert not v["ok"]
    assert {c["id"]: c["status"] for c in v["checks"]}["content"] == "fail"


def test_docx_with_table_fails():
    doc = Document(io.BytesIO(render_resume_docx(TAILORED)))
    doc.add_table(rows=1, cols=2)
    buf = io.BytesIO(); doc.save(buf)
    v = validate_docx(buf.getvalue(), TAILORED)
    assert not v["ok"] and {c["id"]: c["status"] for c in v["checks"]}["layout"] == "fail"


def test_docx_with_header_text_warns_but_is_not_blocked():
    doc = Document(io.BytesIO(render_resume_docx(TAILORED)))
    doc.sections[0].header.paragraphs[0].text = "Rahul Sharma"
    buf = io.BytesIO(); doc.save(buf)
    v = validate_docx(buf.getvalue(), TAILORED)
    assert v["ok"] and {c["id"]: c["status"] for c in v["checks"]}["headerfooter"] == "warn"


def test_three_page_resume_warns():
    long = {"profile": {"name": "A B", "email": "a@b.co"},
            "experience": [{"title": "Role", "content": "\n".join(f"- line {i} " + "word " * 30 for i in range(110))}]}
    v = validate_pdf(render_resume_pdf(long), long)
    assert v["pages"] >= 3 and v["ok"]
    assert {c["id"]: c["status"] for c in v["checks"]}["length"] == "warn"


# ------------------------------------------------------------------ re-scoring after the user edits
def test_rescore_uses_same_terms_and_keeps_before():
    m = compare(ORIGINAL, TAILORED, JD)
    stripped = {**TAILORED, "skills": ["Python"], "summary": "", "experience": [], "projects": []}
    m2 = rescore(m, stripped)
    assert m2["before"] == m["before"]
    assert m2["after"]["score"] < m["after"]["score"]
    assert m2["after"]["total_terms"] == m["after"]["total_terms"]
    assert "Kubernetes" in m2["still_missing"]


def test_unverified_skills_flags_what_the_original_never_mentioned():
    flagged = unverified_skills(ORIGINAL, TAILORED)
    assert "Kubernetes" in flagged and "AWS" in flagged
    assert "Python" not in flagged and "Docker" not in flagged
    assert unverified_skills("", TAILORED) == []          # nothing to compare against -> no false alarms


# ------------------------------------------------------------------ readiness verdict
def test_verdict_is_the_more_cautious_of_score_and_ai():
    m = compare(ORIGINAL, TAILORED, JD)
    m["after"]["score"] = 90
    assert verdict(m, {"overall_readiness": "ready", "gaps": []})["level"] == "ready"
    assert verdict(m, {"overall_readiness": "needs_work", "gaps": []})["level"] == "needs_work"      # AI is stricter
    m["after"]["score"] = 30
    assert verdict(m, {"overall_readiness": "ready", "gaps": []})["level"] == "significant_gap"      # measured match is stricter


def test_high_priority_gap_blocks_ready_and_is_explained():
    m = compare(ORIGINAL, TAILORED, JD); m["after"]["score"] = 90
    v = verdict(m, {"overall_readiness": "ready", "gaps": [{"severity": "high"}]})
    assert v["level"] == "almost" and v["counts"]["high"] == 1
    assert any("high-priority" in f["text"] for f in v["factors"])
    v3 = verdict(m, {"overall_readiness": "ready", "gaps": [{"severity": "high"}] * 3})
    assert v3["level"] == "needs_work"


def test_verdict_without_any_data_is_none():
    assert verdict(None, None) is None


# ------------------------------------------------------------------ templates must actually look different
def _name_x_and_fonts(pdf: bytes):
    """x position of the candidate name and the fonts used on page 1."""
    from pypdf import PdfReader

    page = PdfReader(io.BytesIO(pdf)).pages[0]
    seen = {}

    def visit(text, cm, tm, fd, fs):
        if "Rahul Sharma" in text and "name_x" not in seen:
            seen["name_x"] = tm[4] * cm[0] + cm[4]

    page.extract_text(visitor_text=visit)
    fonts = {str(f.get_object()["/BaseFont"]) for f in page["/Resources"]["/Font"].values()}
    return seen["name_x"], fonts


def test_templates_have_distinct_looks():
    looks = {(t.font, t.align, t.accent, t.head_bar, t.rule) for t in TEMPLATES.values()}
    assert len(looks) == len(TEMPLATES)                                   # no two templates share the same look


def test_pdf_templates_render_differently():
    pdfs = {k: render_resume_pdf(TAILORED, template=k) for k in TEMPLATES}
    x = {k: _name_x_and_fonts(v) for k, v in pdfs.items()}
    assert x["ats_serif"][0] > x["ats_classic"][0] + 100                  # serif: name is centered, classic: left
    assert any("Times" in f for f in x["ats_serif"][1]) and not any("Times" in f for f in x["ats_classic"][1])
    assert len({v for v in pdfs.values()}) == len(pdfs)
    # compact draws shaded heading bars (a fill colour operator in the page content), classic does not
    from pypdf import PdfReader

    import re

    def has_bar(pdf):
        stream = PdfReader(io.BytesIO(pdf)).pages[0].get_contents().get_data().decode("latin-1")
        return re.search(r"\.9333\d* \.9450\d* \.9607\d* rg", stream) is not None     # the #EEF1F5 fill

    assert has_bar(pdfs["ats_compact"]) and not has_bar(pdfs["ats_classic"])


def test_docx_templates_render_differently():
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = {k: Document(io.BytesIO(render_resume_docx(TAILORED, template=k))) for k in TEMPLATES}
    assert d["ats_serif"].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert d["ats_classic"].paragraphs[0].alignment in (None, WD_ALIGN_PARAGRAPH.LEFT)
    heading_xml = lambda doc: next(p for p in doc.paragraphs if p.style.name == "Heading 1")._p.xml
    assert "w:shd" in heading_xml(d["ats_compact"]) and "w:shd" not in heading_xml(d["ats_classic"])
    assert "w:pBdr" in heading_xml(d["ats_classic"]) and "w:pBdr" not in heading_xml(d["ats_compact"])
    assert d["ats_compact"].paragraphs[0].runs[0].font.color.rgb != d["ats_classic"].paragraphs[0].runs[0].font.color.rgb
